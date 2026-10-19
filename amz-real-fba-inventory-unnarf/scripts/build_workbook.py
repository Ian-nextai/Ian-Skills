#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl"]
# ///
"""
Build the amz-real-fba-inventory-unnarf Excel workbook from the JSON produced by
scripts/fetch_real_fba_inventory.mjs.

Workbook sheets, in order:
  1. 汇总        one row per marketplace + methodology notes
  2. one sheet per marketplace (real local FBA inventory only)
  3. 数据检查    report metadata, completeness checks, warnings
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
SUBTITLE_FONT = Font(bold=True, size=11, color="1F4E78")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="FCE4D6")
GOOD_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)

FIXED_DISPOSITIONS = [
    ("SELLABLE", "可售"),
    ("WAREHOUSE_DAMAGED", "仓库损坏"),
    ("CUSTOMER_DAMAGED", "客户损坏"),
    ("CARRIER_DAMAGED", "承运损坏"),
    ("DISTRIBUTOR_DAMAGED", "分销商损坏"),
    ("DEFECTIVE", "残次"),
    ("EXPIRED", "过期"),
]
KNOWN = {code for code, _ in FIXED_DISPOSITIONS} | {"UNKNOWN"}
SITE_ORDER = ["US", "CA", "MX", "BR"]


def ordered_site_codes(sites: dict) -> list[str]:
    preferred = [code for code in SITE_ORDER if code in sites]
    others = sorted(code for code in sites if code not in SITE_ORDER)
    return preferred + others


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the real FBA inventory workbook.")
    parser.add_argument("--input", required=True, help="JSON produced by fetch_real_fba_inventory.mjs")
    parser.add_argument("--output", help="Output .xlsx path")
    return parser.parse_args()


def style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def set_widths(ws, widths: list[float]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def safe_snapshot_token(value: str) -> str:
    return re.sub(r"[^0-9]", "", value) or "unknown"


def note_row(ws, text: str, ncols: int, fill: PatternFill | None = None) -> None:
    row = ws.max_row + 1
    ws.cell(row=row, column=1, value=text)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1)
    cell.alignment = WRAP_TOP
    if fill:
        for col in range(1, ncols + 1):
            ws.cell(row=row, column=col).fill = fill


def build_summary(ws, payload: dict) -> None:
    sites = payload.get("sites", {})
    reports = {report["reportId"]: report for report in payload.get("reports", [])}
    ncols = 12

    ws["A1"] = "Amazon 真实 FBA 库存（排除远程配送 / NARF）"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)

    ws["A2"] = f"数据日期：{payload.get('snapshotDate', '-')}"
    ws["A2"].font = Font(bold=True, size=11)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)

    ws["A3"] = (
        "口径：仅统计物理存放在站点所在国 FBA 仓的库存（Inventory Ledger 的 Location 列）。"
        "NARF / 远程配送共享的美国库存不计入任何非美国站点。"
    )
    ws["A3"].alignment = WRAP_TOP
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    ws.row_dimensions[3].height = 30

    header_row = 5
    headers = [
        "站点",
        "国家/地区",
        "Marketplace ID",
        "台账日期",
        "台账粒度",
        "ASIN 数",
        "结存总数",
        "可售",
        "不可售",
        "可售占比",
        "报告 ID",
        "备注",
    ]
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=index, value=header)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 28

    row = header_row + 1
    total_units = total_sellable = total_non_sellable = total_asins = 0
    for country_code in ordered_site_codes(sites):
        site = sites[country_code]
        report = reports.get(site.get("reportId"), {})
        units = int(site.get("totalUnits") or 0)
        sellable = int(site.get("sellableUnits") or 0)
        non_sellable = int(site.get("nonSellableUnits") or 0)
        asin_count = int(site.get("asinCount") or 0)
        total_units += units
        total_sellable += sellable
        total_non_sellable += non_sellable
        total_asins += asin_count
        ratio = sellable / units if units else 0
        note = "仅本地仓" if asin_count else "台账中无该站点本地仓库存"
        if report.get("latestDateComplete") is False:
            note += "；最新日期分组可能不完整"
        values = [
            site.get("sheetName") or country_code,
            country_code,
            site.get("marketplaceId"),
            site.get("snapshotDate"),
            "日" if site.get("dateKind") == "daily" else ("月" if site.get("dateKind") == "monthly" else "-"),
            asin_count,
            units,
            sellable,
            non_sellable,
            ratio,
            site.get("reportId"),
            note,
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=index == 12)
        ws.cell(row=row, column=7).number_format = "#,##0"
        ws.cell(row=row, column=8).number_format = "#,##0"
        ws.cell(row=row, column=9).number_format = "#,##0"
        ws.cell(row=row, column=10).number_format = "0.0%"
        if non_sellable:
            ws.cell(row=row, column=9).fill = BAD_FILL
        row += 1

    total_values = [
        "合计",
        "",
        "",
        "",
        "",
        total_asins,
        total_units,
        total_sellable,
        total_non_sellable,
        (total_sellable / total_units) if total_units else 0,
        "",
        "所有站点本地仓合计；ASIN 数为各站点之和（未跨站点去重）",
    ]
    for index, value in enumerate(total_values, start=1):
        cell = ws.cell(row=row, column=index, value=value)
        cell.font = Font(bold=True)
        cell.fill = TOTAL_FILL
        cell.border = BORDER
    ws.cell(row=row, column=7).number_format = "#,##0"
    ws.cell(row=row, column=8).number_format = "#,##0"
    ws.cell(row=row, column=9).number_format = "#,##0"
    ws.cell(row=row, column=10).number_format = "0.0%"
    row += 2

    notes = [
        "口径说明",
        "1. 本表只统计物理存放在站点所在国 FBA 仓的库存，按 Inventory Ledger Summary 的 Location 列筛选。",
        "2. NARF（北美远程配送）会让美国仓库存显示为加拿大/墨西哥站可售；本表已将其排除，不使用 FBA Inventory API 或 All Inventory 报告的 fulfillable 数量。",
        "3. 台账日期取所选报告中最新的完整日期分组；Ledger 数据通常有 2–7 天延迟，不是实时库存。",
        "4. 周/月粒度的 Ledger 报告按期末结存计算；本表优先选择日粒度报告。",
        "5. 某个站点在台账中没有该 Location 的行，表示该国 FBA 仓没有库存（通常全部依赖远程配送或本地无货），不是数据缺失。",
        "6. 台账行的 Ending Warehouse Balance 按 ASIN + MSKU + Disposition 汇总；不可售包含仓库损坏、客户损坏、承运损坏、分销商损坏、残次、过期及其他。",
    ]
    for text in notes:
        note_row(ws, text, ncols, WARN_FILL if text == "口径说明" else None)

    set_widths(ws, [16, 12, 18, 12, 10, 9, 11, 11, 11, 10, 14, 32])
    ws.freeze_panes = "A6"


def build_site_sheet(ws, site: dict) -> None:
    headers = [
        "ASIN",
        "MSKU",
        "FNSKU",
        "产品名称",
        "结存总数",
        "可售",
        "不可售合计",
        "仓库损坏",
        "客户损坏",
        "承运损坏",
        "分销商损坏",
        "残次",
        "过期",
        "其他不可售",
        "台账日期",
    ]
    for index, header in enumerate(headers, start=1):
        ws.cell(row=1, column=index, value=header)
    style_header(ws, len(headers))

    items = site.get("items") or []
    row = 2
    for item in items:
        dispositions = item.get("dispositions") or {}
        sellable = float(dispositions.get("SELLABLE") or 0)
        total = float(item.get("quantity") or 0)
        fixed_values = [float(dispositions.get(code) or 0) for code, _ in FIXED_DISPOSITIONS[1:]]
        other = sum(
            float(quantity)
            for code, quantity in dispositions.items()
            if code not in KNOWN
        )
        # UNKNOWN is a known-but-unsellable bucket; keep it in "其他不可售".
        other += float(dispositions.get("UNKNOWN") or 0)
        values = [
            item.get("asin"),
            item.get("msku"),
            item.get("fnsku"),
            item.get("title"),
            total,
            sellable,
            total - sellable,
            *fixed_values,
            other,
            site.get("snapshotDate"),
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=index == 4)
        for index in range(5, 15):
            ws.cell(row=row, column=index).number_format = "#,##0"
        if total - sellable:
            ws.cell(row=row, column=7).fill = BAD_FILL
        row += 1

    if not items:
        note_row(
            ws,
            f"{site.get('snapshotDate')} 台账中没有 Location = {site.get('countryCode')} 的库存记录。"
            "该站点可能没有本地 FBA 仓库存，全部依赖远程配送。",
            len(headers),
            WARN_FILL,
        )

    if items:
        total_row = row
        totals = [
            "合计",
            "",
            "",
            "",
            sum(float(item.get("quantity") or 0) for item in items),
            sum(float((item.get("dispositions") or {}).get("SELLABLE") or 0) for item in items),
            sum(
                float(item.get("quantity") or 0)
                - float((item.get("dispositions") or {}).get("SELLABLE") or 0)
                for item in items
            ),
        ]
        for code, _ in FIXED_DISPOSITIONS[1:]:
            totals.append(
                sum(float((item.get("dispositions") or {}).get(code) or 0) for item in items)
            )
        totals.append(
            sum(
                sum(
                    float(quantity)
                    for name, quantity in (item.get("dispositions") or {}).items()
                    if name not in KNOWN
                )
                + float((item.get("dispositions") or {}).get("UNKNOWN") or 0)
                for item in items
            )
        )
        totals.append("")
        for index, value in enumerate(totals, start=1):
            cell = ws.cell(row=total_row, column=index, value=value)
            cell.font = Font(bold=True)
            cell.fill = TOTAL_FILL
            cell.border = BORDER
        ws.auto_filter.ref = f"A1:O{total_row - 1}"

    set_widths(ws, [14, 22, 14, 60, 10, 10, 10, 10, 10, 10, 12, 8, 8, 12, 12])
    ws.freeze_panes = "A2"


def build_checks(ws, payload: dict) -> None:
    ncols = 6
    ws["A1"] = "数据检查"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)

    rows = [
        ("生成时间", payload.get("generatedAt", "-")),
        ("快照日期", payload.get("snapshotDate", "-")),
        ("数据来源", payload.get("method", {}).get("source", "-")),
    ]
    row = 3
    for label, value in rows:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Ledger 报告").font = SUBTITLE_FONT
    row += 1
    report_headers = [
        "报告 ID",
        "数据范围",
        "总行数",
        "已读行数",
        "是否截断",
        "日期倒序",
        "最新日期",
        "最新日行数",
        "最新日完整",
    ]
    for index, header in enumerate(report_headers, start=1):
        cell = ws.cell(row=row, column=index, value=header)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row += 1
    for report in payload.get("reports", []):
        values = [
            report.get("reportId"),
            f"{report.get('dataStartTime', '')} → {report.get('dataEndTime', '')}",
            report.get("rowCount"),
            report.get("returnedRows"),
            "是" if report.get("truncated") else "否",
            "是" if report.get("orderingVerified") else "否",
            report.get("latestDate"),
            report.get("latestDateRows"),
            "是" if report.get("latestDateComplete") else "否",
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=index == 2)
        if report.get("latestDateComplete") is False or not report.get("orderingVerified"):
            ws.cell(row=row, column=9).fill = BAD_FILL
        else:
            ws.cell(row=row, column=9).fill = GOOD_FILL
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="按 Location 汇总（最新日期）").font = SUBTITLE_FONT
    row += 1
    for report in payload.get("reports", []):
        ws.cell(row=row, column=1, value=f"报告 {report.get('reportId')}")
        row += 1
        for location, units in sorted((report.get("locationTotals") or {}).items()):
            ws.cell(row=row, column=2, value=f"Location {location}：{units} 件 / {report.get('locationRows', {}).get(location, 0)} 行")
            row += 1
        for unmapped in report.get("unmappedLocations") or []:
            ws.cell(row=row, column=2, value=f"未映射 Location {unmapped['location']}：{unmapped['units']} 件")
            ws.cell(row=row, column=2).fill = WARN_FILL
            row += 1

    row += 1
    ws.cell(row=row, column=1, value="站点核对").font = SUBTITLE_FONT
    row += 1
    site_headers = ["站点", "国家", "台账日期", "ASIN 数", "结存总数", "备注"]
    for index, header in enumerate(site_headers, start=1):
        cell = ws.cell(row=row, column=index, value=header)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
    row += 1
    for country_code in ordered_site_codes(payload.get("sites", {})):
        site = payload["sites"][country_code]
        values = [
            site.get("sheetName"),
            country_code,
            site.get("snapshotDate"),
            site.get("asinCount"),
            site.get("totalUnits"),
            "本地仓有库存" if site.get("asinCount") else "无本地仓库存记录",
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.border = BORDER
        if not site.get("asinCount"):
            ws.cell(row=row, column=6).fill = WARN_FILL
        row += 1

    warnings = payload.get("warnings") or []
    if warnings:
        row += 1
        ws.cell(row=row, column=1, value="警告").font = SUBTITLE_FONT
        row += 1
        for warning in warnings:
            ws.cell(row=row, column=1, value=warning).alignment = WRAP_TOP
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
            for col in range(1, ncols + 1):
                ws.cell(row=row, column=col).fill = WARN_FILL
            row += 1

    set_widths(ws, [18, 60, 14, 10, 12, 24])
    ws.freeze_panes = "A3"


def main() -> None:
    args = parse_args()
    payload = json.loads(Path(args.input).read_text(encoding="utf-8"))

    snapshot = payload.get("snapshotDate", "unknown")
    output = Path(args.output) if args.output else Path(
        f"amz_real_fba_inventory_unnarf_{safe_snapshot_token(snapshot)}.xlsx"
    )

    wb = Workbook()
    summary = wb.active
    summary.title = "汇总"
    build_summary(summary, payload)

    used_names = {"汇总", "数据检查"}
    for country_code in ordered_site_codes(payload.get("sites", {})):
        site = payload["sites"][country_code]
        name = site.get("sheetName") or f"站点 {country_code}"
        base = name
        suffix = 2
        while name in used_names or len(name) > 31:
            name = f"{base[:26]} ({suffix})"
            suffix += 1
        used_names.add(name)
        build_site_sheet(wb.create_sheet(name), site)

    build_checks(wb.create_sheet("数据检查"), payload)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(str(output))


if __name__ == "__main__":
    main()
